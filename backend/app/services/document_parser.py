import io
import re
import logging
from typing import Dict, Any

logger = logging.getLogger("knowledge_ai.document_parser")

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF bytes using pypdf."""
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for idx, page in enumerate(reader.pages):
            page_content = page.extract_text() or ""
            if page_content.strip():
                pages_text.append(f"### Page {idx + 1}\n\n{page_content.strip()}")
        return "\n\n---\n\n".join(pages_text) if pages_text else "No readable text found in PDF."
    except Exception as e:
        logger.error(f"Error parsing PDF: {e}")
        return f"Error extracting PDF text: {str(e)}"

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts structured text and headings from DOCX bytes using python-docx."""
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            # Preserve headings if style indicates heading
            style_name = p.style.name.lower() if p.style and p.style.name else ""
            if "heading 1" in style_name:
                paragraphs.append(f"# {text}")
            elif "heading 2" in style_name:
                paragraphs.append(f"## {text}")
            elif "heading 3" in style_name:
                paragraphs.append(f"### {text}")
            elif "list" in style_name or text.startswith("- ") or text.startswith("* "):
                paragraphs.append(f"* {text.lstrip('*- ')}")
            else:
                paragraphs.append(text)
        
        # Also extract table contents
        for table in doc.tables:
            table_rows = []
            for r_idx, row in enumerate(table.rows):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                row_str = "| " + " | ".join(cells) + " |"
                table_rows.append(row_str)
                if r_idx == 0:
                    sep = "| " + " | ".join(["---"] * len(cells)) + " |"
                    table_rows.append(sep)
            if table_rows:
                paragraphs.append("\n" + "\n".join(table_rows) + "\n")

        return "\n\n".join(paragraphs) if paragraphs else "No readable text found in Word document."
    except Exception as e:
        logger.error(f"Error parsing DOCX: {e}")
        return f"Error extracting DOCX text: {str(e)}"

def parse_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Parses an uploaded file (PDF, DOCX, TXT, MD) and returns formatted markdown with metadata.
    """
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    raw_title = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()

    if ext == "pdf":
        body = extract_text_from_pdf(file_bytes)
        doctype = "PDF Document"
    elif ext in ["docx", "doc"]:
        body = extract_text_from_docx(file_bytes)
        doctype = "Word Document"
    elif ext in ["txt", "md", "markdown"]:
        try:
            body = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            body = file_bytes.decode("latin-1", errors="ignore")
        doctype = "Markdown Document" if ext in ["md", "markdown"] else "Text Document"
    else:
        try:
            body = file_bytes.decode("utf-8")
            doctype = "Uploaded Document"
        except Exception:
            body = "Unsupported document format."
            doctype = "Binary File"

    # Construct formatted Markdown document
    markdown_content = f"# {raw_title}\n\n"
    markdown_content += f"> 📄 **Imported from**: `{filename}` ({doctype})\n"
    markdown_content += f"> 📅 **Uploaded**: Auto-processed into workspace\n\n"
    markdown_content += "---\n\n"
    markdown_content += body

    words_count = len(re.findall(r"\w+", body))

    return {
        "title": raw_title,
        "filename": filename,
        "doctype": doctype,
        "content": markdown_content,
        "wordCount": words_count
    }
